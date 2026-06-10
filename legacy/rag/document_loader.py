"""
Document Loader Module for AI Industrial Platform.

Handles loading and chunking of various document formats including
PDF, DOCX, TXT, and CSV files.
"""

import os
from typing import List, Optional, Dict, Any
from pathlib import Path
from dataclasses import dataclass, field
import hashlib
import structlog

logger = structlog.get_logger(__name__)


@dataclass
class Document:
    """
    Represents a document chunk with content and metadata.
    
    Attributes:
        content: The text content of the document/chunk.
        metadata: Dictionary containing document metadata.
        doc_id: Unique identifier for the document.
    """
    content: str
    metadata: Dict[str, Any] = field(default_factory=dict)
    doc_id: Optional[str] = None
    
    def __post_init__(self):
        """Generate doc_id if not provided."""
        if self.doc_id is None:
            # Generate ID from content hash
            content_hash = hashlib.md5(self.content.encode()).hexdigest()[:12]
            source = self.metadata.get("source", "unknown")
            self.doc_id = f"{Path(source).stem}_{content_hash}"


class DocumentLoader:
    """
    Multi-format document loader with chunking support.
    
    Supports loading PDF, DOCX, TXT, and CSV files, with configurable
    chunking strategies for optimal RAG retrieval.
    
    Attributes:
        chunk_size: Target size for text chunks in characters.
        chunk_overlap: Overlap between consecutive chunks.
    """
    
    def __init__(
        self,
        chunk_size: Optional[int] = None,
        chunk_overlap: Optional[int] = None
    ):
        """
        Initialize the document loader.
        
        Args:
            chunk_size: Target chunk size in characters. Defaults to config value.
            chunk_overlap: Overlap between chunks. Defaults to config value.
        """
        from config.settings import get_settings
        settings = get_settings()
        
        # Convert token-based settings to character-based (approx 4 chars per token)
        self.chunk_size = chunk_size or (settings.CHUNK_SIZE * 4)
        self.chunk_overlap = chunk_overlap or (settings.CHUNK_OVERLAP * 4)
        
        self._separators = ["\n\n", "\n", ". ", " ", ""]
    
    def load_pdf(self, file_path: str) -> List[Document]:
        """
        Load a PDF file and extract text content.
        
        Args:
            file_path: Path to the PDF file.
        
        Returns:
            List of Document objects with PDF content.
        """
        try:
            from pypdf import PdfReader
            
            reader = PdfReader(file_path)
            documents = []
            
            for page_num, page in enumerate(reader.pages, start=1):
                text = page.extract_text()
                if text and text.strip():
                    doc = Document(
                        content=text.strip(),
                        metadata={
                            "source": file_path,
                            "file_type": "pdf",
                            "page_number": page_num,
                            "total_pages": len(reader.pages)
                        }
                    )
                    documents.append(doc)
            
            logger.info(
                "PDF loaded",
                path=file_path,
                pages=len(reader.pages),
                documents=len(documents)
            )
            return documents
        
        except ImportError:
            logger.error("pypdf not installed. Run: pip install pypdf")
            raise
        except Exception as e:
            logger.error("PDF loading failed", path=file_path, error=str(e))
            raise
    
    def load_docx(self, file_path: str) -> List[Document]:
        """
        Load a DOCX file and extract text content.
        
        Args:
            file_path: Path to the DOCX file.
        
        Returns:
            List of Document objects with DOCX content.
        """
        try:
            from docx import Document as DocxDocument
            
            doc = DocxDocument(file_path)
            full_text = []
            
            for para in doc.paragraphs:
                if para.text.strip():
                    full_text.append(para.text.strip())
            
            # Also extract text from tables
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(
                        cell.text.strip() for cell in row.cells if cell.text.strip()
                    )
                    if row_text:
                        full_text.append(row_text)
            
            content = "\n\n".join(full_text)
            
            document = Document(
                content=content,
                metadata={
                    "source": file_path,
                    "file_type": "docx",
                    "paragraph_count": len(doc.paragraphs),
                    "table_count": len(doc.tables)
                }
            )
            
            logger.info(
                "DOCX loaded",
                path=file_path,
                paragraphs=len(doc.paragraphs)
            )
            return [document]
        
        except ImportError:
            logger.error("python-docx not installed. Run: pip install python-docx")
            raise
        except Exception as e:
            logger.error("DOCX loading failed", path=file_path, error=str(e))
            raise
    
    def load_txt(self, file_path: str) -> List[Document]:
        """
        Load a TXT file.
        
        Args:
            file_path: Path to the TXT file.
        
        Returns:
            List of Document objects with text content.
        """
        try:
            with open(file_path, "r", encoding="utf-8") as f:
                content = f.read()
            
            document = Document(
                content=content.strip(),
                metadata={
                    "source": file_path,
                    "file_type": "txt",
                    "character_count": len(content)
                }
            )
            
            logger.info("TXT loaded", path=file_path, chars=len(content))
            return [document]
        
        except Exception as e:
            logger.error("TXT loading failed", path=file_path, error=str(e))
            raise
    
    def load_csv(self, file_path: str) -> List[Document]:
        """
        Load a CSV file and convert to text representation.
        
        Args:
            file_path: Path to the CSV file.
        
        Returns:
            List of Document objects with CSV content.
        """
        try:
            import pandas as pd
            
            df = pd.read_csv(file_path)
            
            # Convert DataFrame to readable text
            text_parts = []
            
            # Add column descriptions
            columns_desc = f"Columns: {', '.join(df.columns)}"
            text_parts.append(columns_desc)
            text_parts.append(f"Total rows: {len(df)}")
            text_parts.append("")
            
            # Add data as text (first 100 rows to avoid huge documents)
            sample_df = df.head(100)
            for idx, row in sample_df.iterrows():
                row_text = " | ".join(f"{col}: {val}" for col, val in row.items())
                text_parts.append(row_text)
            
            if len(df) > 100:
                text_parts.append(f"... and {len(df) - 100} more rows")
            
            content = "\n".join(text_parts)
            
            document = Document(
                content=content,
                metadata={
                    "source": file_path,
                    "file_type": "csv",
                    "row_count": len(df),
                    "column_count": len(df.columns),
                    "columns": list(df.columns)
                }
            )
            
            logger.info(
                "CSV loaded",
                path=file_path,
                rows=len(df),
                columns=len(df.columns)
            )
            return [document]
        
        except ImportError:
            logger.error("pandas not installed. Run: pip install pandas")
            raise
        except Exception as e:
            logger.error("CSV loading failed", path=file_path, error=str(e))
            raise
    
    def load_xlsx(self, file_path: str) -> List[Document]:
        """
        Load an Excel file and convert to text representation.
        
        Args:
            file_path: Path to the XLSX file.
        
        Returns:
            List of Document objects with Excel content.
        """
        try:
            import pandas as pd
            
            # Load all sheets
            xlsx = pd.ExcelFile(file_path)
            documents = []
            
            for sheet_name in xlsx.sheet_names:
                df = pd.read_excel(xlsx, sheet_name=sheet_name)
                
                text_parts = [f"Sheet: {sheet_name}"]
                text_parts.append(f"Columns: {', '.join(df.columns)}")
                text_parts.append(f"Total rows: {len(df)}")
                text_parts.append("")
                
                sample_df = df.head(100)
                for idx, row in sample_df.iterrows():
                    row_text = " | ".join(f"{col}: {val}" for col, val in row.items())
                    text_parts.append(row_text)
                
                if len(df) > 100:
                    text_parts.append(f"... and {len(df) - 100} more rows")
                
                content = "\n".join(text_parts)
                
                document = Document(
                    content=content,
                    metadata={
                        "source": file_path,
                        "file_type": "xlsx",
                        "sheet_name": sheet_name,
                        "row_count": len(df),
                        "column_count": len(df.columns)
                    }
                )
                documents.append(document)
            
            logger.info(
                "XLSX loaded",
                path=file_path,
                sheets=len(xlsx.sheet_names)
            )
            return documents
        
        except ImportError:
            logger.error("openpyxl not installed. Run: pip install openpyxl")
            raise
        except Exception as e:
            logger.error("XLSX loading failed", path=file_path, error=str(e))
            raise
    
    def load(self, file_path: str) -> List[Document]:
        """
        Load a file based on its extension.
        
        Args:
            file_path: Path to the file to load.
        
        Returns:
            List of Document objects.
        
        Raises:
            ValueError: If file type is not supported.
        """
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"File not found: {file_path}")
        
        ext = Path(file_path).suffix.lower()
        
        loaders = {
            ".pdf": self.load_pdf,
            ".docx": self.load_docx,
            ".txt": self.load_txt,
            ".csv": self.load_csv,
            ".xlsx": self.load_xlsx,
            ".md": self.load_txt,  # Markdown as text
        }
        
        loader = loaders.get(ext)
        if loader is None:
            raise ValueError(f"Unsupported file type: {ext}")
        
        return loader(file_path)
    
    def _split_text(self, text: str, separators: List[str]) -> List[str]:
        """
        Recursively split text using separators.
        
        Args:
            text: Text to split.
            separators: List of separators to try in order.
        
        Returns:
            List of text chunks.
        """
        final_chunks = []
        
        # Get the appropriate separator
        separator = separators[-1]  # Default to last (empty string)
        new_separators = []
        
        for i, sep in enumerate(separators):
            if sep == "":
                separator = sep
                break
            if sep in text:
                separator = sep
                new_separators = separators[i + 1:]
                break
        
        # Split by the separator
        if separator:
            splits = text.split(separator)
        else:
            splits = list(text)
        
        # Process splits
        good_splits = []
        _separator = "" if separator == "" else separator
        
        for split in splits:
            if len(split) < self.chunk_size:
                good_splits.append(split)
            else:
                if good_splits:
                    merged = self._merge_splits(good_splits, _separator)
                    final_chunks.extend(merged)
                    good_splits = []
                
                if new_separators:
                    other_chunks = self._split_text(split, new_separators)
                    final_chunks.extend(other_chunks)
                else:
                    final_chunks.append(split)
        
        if good_splits:
            merged = self._merge_splits(good_splits, _separator)
            final_chunks.extend(merged)
        
        return final_chunks
    
    def _merge_splits(self, splits: List[str], separator: str) -> List[str]:
        """
        Merge small splits into larger chunks.
        
        Args:
            splits: List of text splits.
            separator: Separator to use when merging.
        
        Returns:
            List of merged chunks.
        """
        merged = []
        current_chunk = []
        current_length = 0
        
        for split in splits:
            split_len = len(split)
            
            if current_length + split_len + len(separator) > self.chunk_size:
                if current_chunk:
                    merged.append(separator.join(current_chunk))
                
                # Handle overlap
                while current_length > self.chunk_overlap and current_chunk:
                    current_length -= len(current_chunk[0]) + len(separator)
                    current_chunk = current_chunk[1:]
            
            current_chunk.append(split)
            current_length += split_len + len(separator)
        
        if current_chunk:
            merged.append(separator.join(current_chunk))
        
        return merged
    
    def chunk_documents(
        self,
        documents: List[Document],
        chunk_size: Optional[int] = None,
        chunk_overlap: Optional[int] = None
    ) -> List[Document]:
        """
        Split documents into smaller chunks for embedding.
        
        Args:
            documents: List of Document objects to chunk.
            chunk_size: Override default chunk size.
            chunk_overlap: Override default overlap.
        
        Returns:
            List of chunked Document objects.
        """
        if chunk_size:
            self.chunk_size = chunk_size
        if chunk_overlap:
            self.chunk_overlap = chunk_overlap
        
        chunked_docs = []
        
        for doc in documents:
            chunks = self._split_text(doc.content, self._separators)
            
            for i, chunk in enumerate(chunks):
                if not chunk.strip():
                    continue
                
                chunk_metadata = doc.metadata.copy()
                chunk_metadata["chunk_index"] = i
                chunk_metadata["total_chunks"] = len(chunks)
                
                chunked_doc = Document(
                    content=chunk.strip(),
                    metadata=chunk_metadata
                )
                chunked_docs.append(chunked_doc)
        
        logger.info(
            "Documents chunked",
            input_docs=len(documents),
            output_chunks=len(chunked_docs),
            chunk_size=self.chunk_size,
            overlap=self.chunk_overlap
        )
        
        return chunked_docs
    
    def load_and_chunk(
        self,
        file_path: str,
        department: Optional[str] = None
    ) -> List[Document]:
        """
        Load a file and chunk it in one step.
        
        Args:
            file_path: Path to the file to load.
            department: Optional department to add to metadata.
        
        Returns:
            List of chunked Document objects.
        """
        documents = self.load(file_path)
        
        # Add department to metadata
        if department:
            for doc in documents:
                doc.metadata["department"] = department
        
        return self.chunk_documents(documents)


if __name__ == "__main__":
    # Test document loader
    loader = DocumentLoader(chunk_size=500, chunk_overlap=50)
    
    # Create a test text file
    test_content = """
    Industrial Maintenance Manual
    
    Chapter 1: Safety Procedures
    
    Before performing any maintenance work, ensure all safety protocols are followed.
    Lock out tag out (LOTO) procedures must be strictly adhered to.
    Personal protective equipment (PPE) including safety glasses, gloves, and 
    steel-toed boots are mandatory in all maintenance areas.
    
    Chapter 2: Preventive Maintenance
    
    Regular preventive maintenance is essential for equipment longevity.
    Schedule inspections based on manufacturer recommendations and operational hours.
    Document all maintenance activities in the maintenance management system.
    
    Chapter 3: Troubleshooting
    
    When equipment failures occur, follow the systematic troubleshooting approach:
    1. Identify the symptoms
    2. Review recent changes or events
    3. Check common failure points
    4. Use diagnostic tools
    5. Document findings and resolution
    """ * 3  # Repeat to make longer
    
    test_file = "test_document.txt"
    with open(test_file, "w") as f:
        f.write(test_content)
    
    # Load and chunk
    print("Loading document...")
    docs = loader.load(test_file)
    print(f"Loaded {len(docs)} document(s)")
    
    print("\nChunking documents...")
    chunks = loader.chunk_documents(docs)
    print(f"Created {len(chunks)} chunks")
    
    for i, chunk in enumerate(chunks[:3]):
        print(f"\nChunk {i + 1}:")
        print(f"  Length: {len(chunk.content)}")
        print(f"  Preview: {chunk.content[:100]}...")
        print(f"  Metadata: {chunk.metadata}")
    
    # Cleanup
    import os
    os.remove(test_file)
    print("\nTest completed!")
